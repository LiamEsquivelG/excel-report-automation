"""
excel-report-automation
-----------------------
Automatiza un reporte mensual recurrente: toma datos crudos, los limpia,
calcula KPIs y genera un Excel formateado con varias hojas y un grafico.

Inspirado en tareas que automatice en mis practicas (finanzas y administracion).
Datos 100% SIMULADOS.

Uso:
  python report_automation.py --month 2025-06 --out reporte_2025-06.xlsx
"""
from __future__ import annotations

import argparse
import time

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AREAS = ["Operaciones", "Finanzas", "Comercial", "RR.HH.", "TI"]
CATEGORIES = ["Servicios externos", "Insumos", "Viajes", "Software", "Capacitacion"]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def generate_raw(month: str, rng: np.random.Generator, n: int = 800) -> pd.DataFrame:
    """Simula un export 'sucio' de un ERP: duplicados, montos como texto, areas mal escritas."""
    start = pd.Period(month).start_time
    days = pd.Period(month).days_in_month
    df = pd.DataFrame({
        "id_doc": rng.integers(10_000, 99_999, n),
        "fecha": start + pd.to_timedelta(rng.integers(0, days, n), unit="D"),
        "area": rng.choice(AREAS + ["operaciones ", "TI "], n),
        "categoria": rng.choice(CATEGORIES, n),
        "monto": rng.gamma(2.0, 350_000, n).round(-2),
        "presupuesto": 0.0,
    })
    df["monto"] = df.monto.map(lambda x: f"{x:,.0f}".replace(",", "."))  # "1.234.500"
    return pd.concat([df, df.sample(25, random_state=1)])                  # duplicados


def clean(raw: pd.DataFrame) -> pd.DataFrame:
    df = raw.drop_duplicates().copy()
    df["area"] = df.area.str.strip().str.capitalize().replace({"Ti": "TI", "Rr.hh.": "RR.HH."})
    df["monto"] = df.monto.str.replace(".", "", regex=False).astype(float)
    return df.drop(columns="presupuesto")


def summarize(df: pd.DataFrame, rng: np.random.Generator) -> dict[str, pd.DataFrame]:
    by_area = df.groupby("area", as_index=False).agg(gasto=("monto", "sum"), documentos=("id_doc", "count"))
    by_area["presupuesto"] = (by_area.gasto * rng.uniform(0.9, 1.15, len(by_area))).round(-4)
    by_area["desviacion_%"] = (100 * (by_area.gasto / by_area.presupuesto - 1)).round(1)
    by_area["estado"] = np.where(by_area["desviacion_%"] > 5, "Sobre presupuesto", "OK")
    pivot = df.pivot_table(index="area", columns="categoria", values="monto", aggfunc="sum", fill_value=0)
    weekly = df.groupby(df.fecha.dt.isocalendar().week)["monto"].sum().rename_axis("semana").reset_index()
    kpis = pd.DataFrame({
        "KPI": ["Gasto total", "N° documentos", "Ticket promedio", "Areas sobre presupuesto"],
        "Valor": [df.monto.sum(), len(df), df.monto.mean(), int((by_area.estado != "OK").sum())],
    })
    return {"Resumen": kpis, "Por area": by_area, "Area x categoria": pivot.reset_index(),
            "Semanal": weekly, "Detalle": df.sort_values("fecha")}


def write_excel(sheets: dict[str, pd.DataFrame], path: str) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, data in sheets.items():
            data.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 3, 40)
            for c in col[1:]:
                if isinstance(c.value, (int, float)) and abs(c.value) >= 1000:
                    c.number_format = "#,##0"
        ws.freeze_panes = "A2"
    ws = wb["Por area"]
    chart = BarChart()
    chart.title, chart.y_axis.title = "Gasto vs presupuesto por area", "CLP"
    chart.add_data(Reference(ws, min_col=2, max_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.add_data(Reference(ws, min_col=4, max_col=4, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    ws.add_chart(chart, "H2")
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", default="2025-06")
    ap.add_argument("--out", default=None)
    ap.add_argument("--seed", type=int, default=11)
    args = ap.parse_args()
    out = args.out or f"reporte_{args.month}.xlsx"

    t0 = time.perf_counter()
    rng = np.random.default_rng(args.seed)
    raw = generate_raw(args.month, rng)
    df = clean(raw)
    sheets = summarize(df, rng)
    write_excel(sheets, out)
    elapsed = time.perf_counter() - t0

    print(f"Filas crudas: {len(raw)} -> limpias: {len(df)} (duplicados eliminados: {len(raw) - len(df)})")
    print(sheets["Por area"].to_string(index=False))
    print(f"\nReporte generado: {out} en {elapsed:.1f} s")


if __name__ == "__main__":
    main()
